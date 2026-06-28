"""Non-régression extraction de texte (formats clés, sans dépendance lourde)."""

import io

from docx import Document  # type: ignore

from app.services.file_extractor import (
    _rows_to_markdown_table,
    detect_source_type,
    extract_text_from_bytes,
)
from app.services.chunker import chunk_text_hierarchical


class TestExtractTextFromBytes:
    def test_markdown_by_content_type(self):
        out = extract_text_from_bytes(b"# Titre\ncorps", "x.md", "text/markdown")
        assert "Titre" in out and "corps" in out

    def test_plain_text_by_extension(self):
        out = extract_text_from_bytes(b"hello world", "notes.txt", None)
        assert out == "hello world"

    def test_unknown_type_falls_back_to_utf8_decode(self):
        out = extract_text_from_bytes(b"raw bytes", "blob.bin", "application/x-thing")
        assert out == "raw bytes"

    def test_invalid_utf8_is_lenient(self):
        # errors="ignore" : pas d'exception sur des octets non décodables.
        out = extract_text_from_bytes(b"ok\xff\xfe", "x.txt", "text/plain")
        assert "ok" in out


class TestTableExtraction:
    def test_rows_to_markdown_table(self):
        md = _rows_to_markdown_table([["A", "B"], ["1", "2"]])
        assert "| A | B |" in md
        assert "| --- | --- |" in md
        assert "| 1 | 2 |" in md

    def test_docx_includes_table_cells(self):
        doc = Document()
        doc.add_paragraph("Introduction")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Recall@10"
        table.cell(1, 1).text = "0.85"
        buf = io.BytesIO()
        doc.save(buf)

        out = extract_text_from_bytes(
            buf.getvalue(),
            "report.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        assert "Introduction" in out
        assert "Recall@10" in out
        assert "0.85" in out
        assert "| Metric | Value |" in out

    def test_docx_preserves_block_order(self):
        doc = Document()
        doc.add_paragraph("Before table")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Cell"
        doc.add_paragraph("After table")
        buf = io.BytesIO()
        doc.save(buf)

        out = extract_text_from_bytes(
            buf.getvalue(),
            "ordered.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        assert out.index("Before table") < out.index("Cell") < out.index("After table")

    def test_chunker_splits_oversized_markdown_table(self):
        header = "| Col1 | Col2 |"
        sep = "| --- | --- |"
        rows = "\n".join(f"| r{i} | v{i} |" for i in range(80))
        text = f"Intro\n\n{header}\n{sep}\n{rows}\n\nFin"
        chunks = chunk_text_hierarchical(text, chunk_size=50, overlap=10)
        table_chunks = [c for c in chunks if "| Col1 |" in c["text"]]
        assert len(table_chunks) >= 2
        for chunk in table_chunks:
            assert "| --- | --- |" in chunk["text"]

    def test_chunker_preserves_markdown_table_block(self):
        text = "Intro\n\n| Col1 | Col2 |\n| --- | --- |\n| a | b |\n\nFin"
        chunks = chunk_text_hierarchical(text, chunk_size=400, overlap=60)
        table_chunks = [c for c in chunks if "| Col1 |" in c["text"]]
        assert table_chunks
        assert "| a | b |" in table_chunks[0]["text"]


class TestDetectSourceType:
    def test_uses_extension(self):
        assert detect_source_type("doc.PDF", None) == "pdf"

    def test_falls_back_to_content_type(self):
        assert detect_source_type("noext", "text/markdown") == "markdown"

    def test_unknown(self):
        assert detect_source_type("noext", None) == "unknown"
